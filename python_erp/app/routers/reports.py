from fastapi import APIRouter, HTTPException, Depends, Query, Response
from typing import Optional, List, Dict, Any
from datetime import datetime
import io
from app.database import get_supabase_client
from app.security import get_current_user

router = APIRouter(prefix="/reports", tags=["Reports & Analytics"])

@router.get("/transactions")
async def list_transactions(
    student_id: Optional[str] = None,
    service_id: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    limit: int = Query(500, le=5000),
    user: dict = Depends(get_current_user)
):
    client = get_supabase_client()
    if not client:
        return []

    query = client.table("transactions").select("*")
    if student_id:
        query = query.eq("student_id", student_id)
    if service_id:
        query = query.eq("service_id", service_id)
    if start_date:
        query = query.gte("created_at", start_date)
    if end_date:
        query = query.lte("created_at", end_date)

    res = query.order("created_at", desc=True).limit(limit).execute()
    return res.data or []

@router.get("/summary")
async def get_dashboard_summary(user: dict = Depends(get_current_user)):
    client = get_supabase_client()
    if not client:
        return {"total_students": 0, "today_transactions": 0, "today_amount": 0.0}

    students_cnt = client.table("students").select("id", count="exact").execute()
    
    today_iso = datetime.now().strftime("%Y-%m-%dT00:00:00")
    txs = client.table("transactions").select("amount").gte("created_at", today_iso).execute()
    
    today_amount = sum(float(tx.get("amount", 0.0)) for tx in (txs.data or []))

    return {
        "total_students": students_cnt.count or 0,
        "today_transactions": len(txs.data or []),
        "today_amount": today_amount
    }

@router.get("/export-excel")
async def export_transactions_excel(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    user: dict = Depends(get_current_user)
):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed on server.")

    client = get_supabase_client()
    if not client:
        raise HTTPException(status_code=503, detail="Database unavailable")

    query = client.table("transactions").select("*")
    if start_date:
        query = query.gte("created_at", start_date)
    if end_date:
        query = query.lte("created_at", end_date)

    res = query.order("created_at", desc=True).limit(5000).execute()
    records = res.data or []

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gurukul Transactions"

    headers = ["Receipt No", "Date & Time", "SUID", "Student Name", "NFC Card", "Service", "Amount (Rs)"]
    ws.append(headers)

    # Style Header
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r in records:
        ws.append([
            r.get("receipt_no"),
            r.get("created_at", "")[:19].replace("T", " "),
            r.get("suid"),
            r.get("student_name"),
            r.get("nfc_no"),
            r.get("service_name"),
            float(r.get("amount", 0.0))
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"gurukul_transactions_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
